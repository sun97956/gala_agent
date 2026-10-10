# tools.py
"""
GAIA agent tools. All free, no extra API keys needed:
  - web search: ddgs (DuckDuckGo / Bing / Brave ... auto backend)
  - webpages / PDFs: requests + BeautifulSoup + markdownify / pypdf
  - Wikipedia: MediaWiki API, full article incl. tables, optional historical revision
  - audio: faster-whisper (runs locally on CPU)
  - images / video frames: DeepSeek vision (deepseek-flash, same DEEPSEEK_API_KEY)
  - YouTube: youtube-transcript-api, yt-dlp + opencv for frames
  - files: Excel / CSV / PDF / text / Python (and executing the Python file)
  - chess: python-chess (+ Stockfish if installed)
"""

import base64
import io
import json
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from concurrent.futures import TimeoutError as FuturesTimeout
from pathlib import Path

import requests
from smolagents import tool

# =====================================================
# Config
# =====================================================

BASE_DIR = Path(__file__).resolve().parent
ATTACHMENTS_DIR = (BASE_DIR / "attachments").resolve()
CACHE_DIR = (BASE_DIR / ".cache").resolve()

API_BASE = os.getenv("DEEPSEEK_API_BASE", "https://api.deepseek.com")
VISION_MODEL = os.getenv("VISION_MODEL", "deepseek-flash")
WHISPER_MODEL = os.getenv("WHISPER_MODEL", "small")  # tiny / base / small / medium

# Wikimedia rejects generic user agents (HTTP 429/403) -> must be descriptive.
USER_AGENT = (
    "GAIA-Agent-Course/1.0 (educational project; "
    "https://huggingface.co/learn/agents-course)"
)
BROWSER_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
)

# smolagents cuts any printed output above 20000 chars OUT OF THE MIDDLE (silently losing data),
# so tool outputs are paginated below that limit; the agent continues with start=...
MAX_CHARS = 15000

# If HTTPS_PROXY is set (needed for YouTube), don't send DeepSeek / HF-mirror traffic through it:
# a local proxy that can't reach them causes "Connection error" on every model call.
if os.getenv("HTTPS_PROXY") or os.getenv("HTTP_PROXY") or os.getenv("https_proxy"):
    _direct = ["api.deepseek.com", "hf-mirror.com", "localhost", "127.0.0.1"]
    _cur = [x.strip() for x in (os.getenv("NO_PROXY") or os.getenv("no_proxy") or "").split(",") if x.strip()]
    os.environ["NO_PROXY"] = os.environ["no_proxy"] = ",".join(dict.fromkeys(_cur + _direct))

_session = requests.Session()
_session.headers.update({"User-Agent": USER_AGENT})


def _run_with_timeout(func, timeout):
    """Run func() in a thread; raise TimeoutError if it takes too long."""
    ex = ThreadPoolExecutor(max_workers=1)
    fut = ex.submit(func)
    try:
        return fut.result(timeout=timeout)
    except FuturesTimeout:
        raise TimeoutError(f"timed out after {timeout}s")
    finally:
        ex.shutdown(wait=False, cancel_futures=True)


def _clip(text, start=0, max_chars=MAX_CHARS):
    text = text or ""
    total = len(text)
    chunk = text[start:start + max_chars]
    if start + max_chars < total:
        chunk += (
            f"\n\n[... truncated: showing chars {start}-{start + max_chars} of {total}. "
            f"Call again with start={start + max_chars} to read more.]"
        )
    return chunk


def resolve_file(file_name):
    """Accept a bare attachment name or a path; return an existing Path."""
    p = Path(file_name)
    candidates = [p] if p.is_absolute() else [ATTACHMENTS_DIR / p, BASE_DIR / p, p]
    for c in candidates:
        if c.exists():
            return c.resolve()
    raise FileNotFoundError(f"File not found: {file_name} (looked in {ATTACHMENTS_DIR})")


def _get_with_retry(url, params=None, headers=None, timeout=20, tries=4):
    last = None
    for i in range(tries):
        try:
            r = _session.get(url, params=params, headers=headers, timeout=timeout)
            if r.status_code == 429:
                last = "HTTP 429 rate limited"
                wait = r.headers.get("Retry-After", "")
                time.sleep(min(int(wait), 20) if wait.isdigit() else 3 * (i + 1))
                continue
            r.raise_for_status()
            return r
        except requests.RequestException as e:
            last = str(e)
            time.sleep(1 + i)
    raise RuntimeError(f"GET failed for {url}: {last}")


def _html_to_markdown(html, base_url=None):
    from bs4 import BeautifulSoup
    from markdownify import markdownify

    soup = BeautifulSoup(html, "lxml") if _has_lxml() else BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "footer", "nav", "form", "iframe"]):
        tag.decompose()
    # Wikipedia clutter
    # (no CSS selectors: some soupsieve versions on Windows raise SelectorSyntaxError)
    clutter = {"mw-editsection", "reference", "navbox", "metadata", "mw-jump-link"}
    for t in soup.find_all(class_=lambda c: c is not None and bool(clutter & set(str(c).split()))):
        try:
            t.decompose()
        except Exception:
            pass
    md = markdownify(str(soup), heading_style="ATX", strip=["img"])
    md = re.sub(r"\n{3,}", "\n\n", md)
    md = re.sub(r"[ \t]+\n", "\n", md)
    return md.strip()


def _has_lxml():
    try:
        import lxml  # noqa: F401
        return True
    except ImportError:
        return False


# =====================================================
# Attachment download (used by agent.py, not exposed as a tool)
# =====================================================

SCORING_API = "https://agents-course-unit4-scoring.hf.space"


def ensure_attachment(task_id, file_name):
    """
    Return local path of the attachment, downloading it if needed.
    Order: attachments/ folder -> scoring API /files/{task_id}
           -> HF dataset gaia-benchmark/GAIA (needs HF_TOKEN + accepted terms).
    """
    if not file_name:
        return None
    ATTACHMENTS_DIR.mkdir(parents=True, exist_ok=True)
    local = ATTACHMENTS_DIR / file_name
    if local.exists() and local.stat().st_size > 0:
        return local

    try:
        r = requests.get(f"{SCORING_API}/files/{task_id}", timeout=60)
        if r.status_code == 200 and "application/json" not in r.headers.get("content-type", ""):
            local.write_bytes(r.content)
            return local
    except requests.RequestException:
        pass

    try:
        from huggingface_hub import hf_hub_download

        for split in ("validation", "test"):
            try:
                p = hf_hub_download(
                    repo_id="gaia-benchmark/GAIA",
                    repo_type="dataset",
                    filename=f"2023/{split}/{file_name}",
                    token=os.getenv("HF_TOKEN"),
                )
                shutil.copy(p, local)
                return local
            except Exception:
                continue
    except ImportError:
        pass

    print(f"[tools] WARNING: could not obtain attachment {file_name}")
    return None


# =====================================================
# Web search
# =====================================================

# Pages that just echo other people's search queries, or that publish GAIA questions/answers.
# Using them would be answer leakage, not research.
_BLOCKED_URL_PATTERNS = [
    r"gaia-benchmark", r"gaia[_-]?(validation|level|test)", r"_gaia/", r"traces/.*gaia", r"datasets/gaia",
    r"harbor-datasets", r"solve\.sh",
    r"agents-course", r"unit4",
    r"yandex\.[a-z]+/search", r"wordplays\.com", r"fukgames\.com", r"etsy\.com/market",
    r"kidvideo\.org/search",
    # SEO pages generated from other people's search queries (they echo questions + guessed answers)
    r"/search[/?]", r"/tag/", r"instagram\.com/popular/", r"lespac\.com", r"edurev\.in/search",
    # benchmark trace / task-id dumps
    r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}",
]


CURRENT_QUESTION = ""  # set by agent.py for each task, used to spot pages that copy the question


def _words(t):
    return re.findall(r"[a-z0-9]+", (t or "").lower())


def _echoes_question(text, n=8):
    q = _words(CURRENT_QUESTION)
    if len(q) < n:
        return False
    t = " " + " ".join(_words(text)) + " "
    return any(" " + " ".join(q[i:i + n]) + " " in t for i in range(len(q) - n + 1))


def _is_leak_or_spam(r):
    if not str(r.get("href", "")).startswith("http"):
        return True  # ad redirect / relative junk link
    url_title = f"{r.get('href', '')} {r.get('title', '')}".lower()
    body = (r.get("body") or "").lower()
    if any(re.search(p, url_title) for p in _BLOCKED_URL_PATTERNS):
        return True
    # datasets / jsonl dumps that redistribute benchmark questions
    if "huggingface.co/datasets/" in url_title or ".jsonl" in url_title:
        return True
    if "gaia" in url_title and not re.search(r"esa|spacecraft|astrometr|satellite|mission|hypothesis", url_title):
        return True
    if "gaia benchmark" in body or "gaia-benchmark" in body or "task_id" in body:
        return True
    # a page that contains the question itself is a copy of the benchmark (or a query echo)
    return _echoes_question(f"{r.get('title', '')} {body}")


@tool
def web_search(query: str, max_results: int = 8) -> str:
    """
    Search the web (DuckDuckGo/Bing/Brave via ddgs). Returns titles, URLs and snippets.
    Snippets are only leads: open the most relevant URL with visit_webpage to read the facts.

    Args:
        query: Focused search keywords. Put exact names in quotes, add site: filters when useful.
        max_results: Number of results (default 8).
    """
    from ddgs import DDGS

    if re.search(r"\bgaia\b|[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}|answer key|benchmark answer", query, re.I) \
            and not re.search(r"esa|spacecraft|astrometr|satellite|hypothesis", query, re.I):
        return ("Refused: looking up benchmark answers or task ids is not allowed. "
                "Search for the underlying facts (primary sources) instead.")

    last_err = None
    for attempt in range(3):
        try:
            results = _run_with_timeout(
                lambda: list(DDGS(timeout=15).text(query, max_results=max_results, backend="auto")),
                30,
            )
            if not results:
                return f"No results for: {query}. Try different keywords."
            results = [r for r in results if not _is_leak_or_spam(r)]
            if not results:
                return f"No usable results for: {query}. Try different keywords."
            out = []
            for i, r in enumerate(results, 1):
                out.append(f"[{i}] {r.get('title')}\nURL: {r.get('href')}\n{r.get('body')}")
            return "\n\n".join(out)
        except Exception as e:
            last_err = e
            time.sleep(2 + 2 * attempt)
    return f"Search failed ({type(last_err).__name__}: {last_err}). Try a different query or use wikipedia_search."


# =====================================================
# Webpage / PDF
# =====================================================

@tool
def visit_webpage(url: str, start: int = 0) -> str:
    """
    Open a URL and return its main content as markdown (tables preserved). Also handles PDF links.
    Long pages are paginated: use `start` to continue reading.

    Args:
        url: Full http(s) URL.
        start: Character offset to start reading from (default 0).
    """
    if url.startswith("http") and _is_leak_or_spam({"href": url}):
        return f"Blocked: {url} publishes benchmark answers or echoes search queries. Use primary sources."
    try:
        r = _session.get(url, headers={"User-Agent": BROWSER_UA}, timeout=30, allow_redirects=True)
        if r.status_code == 429 or r.status_code == 403:
            r = _session.get(url, headers={"User-Agent": USER_AGENT}, timeout=30)
        r.raise_for_status()
    except Exception as e:
        return f"Failed to open {url}: {type(e).__name__}: {e}"

    ctype = r.headers.get("content-type", "").lower()
    if "pdf" in ctype or url.lower().split("?")[0].endswith(".pdf"):
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(r.content))
            text = "\n".join(
                f"--- Page {i + 1} ---\n{(p.extract_text() or '')}" for i, p in enumerate(reader.pages[:80])
            )
            return _clip(text, start)
        except Exception as e:
            return f"PDF parse failed: {e}"

    if "html" not in ctype and "xml" not in ctype:
        return _clip(r.text, start)

    md = _html_to_markdown(r.text, url)
    return f"URL: {r.url}\n\n" + _clip(md, start)


# =====================================================
# Wikipedia
# =====================================================

WIKI_API = "https://en.wikipedia.org/w/api.php"


@tool
def wikipedia_search(query: str) -> str:
    """
    Search English Wikipedia. Returns the top matching article titles with snippets.
    Then call wikipedia_page(title) to read the full article.

    Args:
        query: Search keywords.
    """
    try:
        r = _get_with_retry(
            WIKI_API,
            params={"action": "query", "list": "search", "srsearch": query,
                    "srlimit": 8, "format": "json"},
        )
        hits = r.json().get("query", {}).get("search", [])
    except Exception as e:
        fallback = web_search(f"site:en.wikipedia.org {query}", 8)
        return f"(Wikipedia API unavailable: {e}; web results instead)\n{fallback}"
    if not hits:
        return f"No Wikipedia results for: {query}"
    out = []
    for h in hits:
        snippet = re.sub(r"<[^>]+>", "", h.get("snippet", ""))
        out.append(f"- {h['title']}: {snippet}")
    return "\n".join(out)


@tool
def wikipedia_page(title: str, as_of_date: str = "", start: int = 0) -> str:
    """
    Read a FULL English Wikipedia article as markdown, including tables (discographies, lists, results...).
    Optionally read the version of the article as it was on a given date
    (e.g. a question says "use the 2022 version of Wikipedia" -> as_of_date="2022-12-31").

    Args:
        title: Exact article title (e.g. "Mercedes Sosa discography" or "Mercedes Sosa").
        as_of_date: Optional "YYYY-MM-DD"; returns the latest revision on or before that date.
        start: Character offset for long articles (default 0).
    """
    try:
        params = {"action": "parse", "prop": "text", "format": "json",
                  "formatversion": 2, "redirects": 1}
        if as_of_date:
            rv = _get_with_retry(WIKI_API, params={
                "action": "query", "prop": "revisions", "titles": title, "rvlimit": 1,
                "rvdir": "older", "rvstart": f"{as_of_date}T23:59:59Z", "rvprop": "ids|timestamp",
                "format": "json", "formatversion": 2, "redirects": 1,
            }).json()
            pages = rv.get("query", {}).get("pages", [])
            revs = pages[0].get("revisions") if pages else None
            if not revs:
                return f"No revision of '{title}' found on or before {as_of_date}."
            params["oldid"] = revs[0]["revid"]
            header = f"Wikipedia: {pages[0]['title']} (revision {revs[0]['revid']} of {revs[0]['timestamp']})"
        else:
            params["page"] = title
            header = f"Wikipedia: {title}"

        data = _get_with_retry(WIKI_API, params=params).json()
        if "error" in data:
            return f"Wikipedia error: {data['error'].get('info')}. Use wikipedia_search to find the exact title."
        html = data["parse"]["text"]
        md = _html_to_markdown(html)
        md = re.sub(r'\[([^\]]*)\]\((?:/wiki/|/w/index\.php)[^)]*\)', r'\1', md)  # drop internal wiki links
        return f"{header}\nURL: https://en.wikipedia.org/wiki/{data['parse']['title'].replace(' ', '_')}\n\n" + _clip(md, start)
    except Exception as e:
        # API rate-limited -> fall back to the normal article URL (different rate-limit bucket)
        try:
            url = f"https://en.wikipedia.org/w/index.php?title={requests.utils.quote(title.replace(' ', '_'))}"
            if "oldid" in locals().get("params", {}):
                url += f"&oldid={params['oldid']}"
            r = _get_with_retry(url)
            return f"Wikipedia (html fallback): {title}\nURL: {r.url}\n\n" + _clip(_html_to_markdown(r.text), start)
        except Exception as e2:
            return f"wikipedia_page failed: {type(e).__name__}: {e}; fallback: {e2}"


# =====================================================
# Files
# =====================================================

TEXT_EXT = {".txt", ".md", ".json", ".jsonl", ".csv", ".tsv", ".py", ".js", ".html", ".xml", ".yaml", ".yml", ".log"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".flac", ".ogg", ".aac"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}


@tool
def read_file(file_name: str, start: int = 0) -> str:
    """
    Read an attached file and return its content as text.
    Supports: Excel (.xlsx/.xls, every sheet and every row), CSV, PDF, Word (.docx), PowerPoint (.pptx),
    and plain-text formats (txt, md, json, py, ...). For audio use transcribe_audio, for images use analyze_image.

    Args:
        file_name: Attachment file name (or path).
        start: Character offset for long files (default 0).
    """
    try:
        path = resolve_file(file_name)
    except FileNotFoundError as e:
        return str(e)
    ext = path.suffix.lower()

    try:
        if ext in {".xlsx", ".xlsm", ".xls"}:
            import pandas as pd

            sheets = pd.read_excel(path, sheet_name=None)
            parts = []
            for name, df in sheets.items():
                parts.append(f"## Sheet: {name}  (shape {df.shape[0]} rows x {df.shape[1]} cols)\n"
                             f"Columns: {list(df.columns)}\n{df.to_csv(index=False)}")
            parts.append(f"(Tip: load with pandas.read_excel(r'{path}', sheet_name=None) to compute exactly.)")
            return _clip("\n\n".join(parts), start)

        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            text = "\n".join(f"--- Page {i + 1} ---\n{(p.extract_text() or '')}" for i, p in enumerate(reader.pages))
            return _clip(text, start)

        if ext == ".docx":
            import docx  # python-docx

            d = docx.Document(str(path))
            text = "\n".join(p.text for p in d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    text += "\n" + " | ".join(c.text for c in row.cells)
            return _clip(text, start)

        if ext == ".pptx":
            from pptx import Presentation

            prs = Presentation(str(path))
            out = []
            for i, s in enumerate(prs.slides, 1):
                texts = [sh.text for sh in s.shapes if getattr(sh, "has_text_frame", False) and sh.text]
                out.append(f"--- Slide {i} ---\n" + "\n".join(texts))
            return _clip("\n".join(out), start)

        if ext in AUDIO_EXT:
            return "This is an audio file. Use transcribe_audio(file_name)."
        if ext in IMAGE_EXT:
            return "This is an image. Use analyze_image(file_name, question)."

        return _clip(path.read_text(encoding="utf-8", errors="replace"), start)
    except Exception as e:
        return f"read_file failed: {type(e).__name__}: {e}"


@tool
def run_python_file(file_name: str, timeout: int = 60) -> str:
    """
    Execute an attached Python file in a separate process and return its stdout/stderr.
    Use this to get the real output of attached code instead of simulating it in your head.

    Args:
        file_name: Attachment file name (or path) of the .py file.
        timeout: Max seconds to run (default 60).
    """
    try:
        path = resolve_file(file_name)
        proc = subprocess.run(
            [sys.executable, str(path)],
            cwd=str(path.parent), capture_output=True, text=True, timeout=timeout,
        )
        return (f"exit code: {proc.returncode}\n--- stdout ---\n{proc.stdout[-8000:]}"
                f"\n--- stderr ---\n{proc.stderr[-4000:]}")
    except subprocess.TimeoutExpired as e:
        out = (e.stdout or b"")
        out = out.decode(errors="replace") if isinstance(out, bytes) else out
        return f"Timed out after {timeout}s. Partial stdout:\n{out[-4000:]}"
    except Exception as e:
        return f"run_python_file failed: {type(e).__name__}: {e}"


# =====================================================
# Audio
# =====================================================

_whisper = None


def _decode_audio_16k(path):
    import av
    import numpy as np

    chunks = []
    with av.open(str(path)) as container:
        resampler = av.AudioResampler(format="s16", layout="mono", rate=16000)
        for frame in container.decode(audio=0):
            for f in resampler.resample(frame) or []:
                chunks.append(f.to_ndarray().reshape(-1))
        for f in resampler.resample(None) or []:
            chunks.append(f.to_ndarray().reshape(-1))
    return np.concatenate(chunks).astype(np.float32) / 32768.0


def transcribe_audio_path(path):
    """Runs IN-PROCESS. Called by the child process started from transcribe_audio()."""
    global _whisper
    from faster_whisper import WhisperModel

    errors = []
    for compute_type in ("int8", "float32"):
        try:
            if _whisper is None:
                print(f"[whisper] loading '{WHISPER_MODEL}' ({compute_type})...", file=sys.stderr, flush=True)
                _whisper = WhisperModel(WHISPER_MODEL, device="cpu", compute_type=compute_type, cpu_threads=4)
            try:
                audio = _decode_audio_16k(path)
                print(f"[whisper] decoded {len(audio) / 16000:.1f}s of audio", file=sys.stderr, flush=True)
            except Exception as e:
                print(f"[whisper] own decoder failed ({e}); using built-in", file=sys.stderr, flush=True)
                audio = str(path)
            segments, info = _whisper.transcribe(audio, beam_size=5, vad_filter=False)
            lines = [f"[{s.start:6.1f}s] {s.text.strip()}" for s in segments]
            return f"(language: {info.language})\n" + "\n".join(lines)
        except Exception as e:
            errors.append(f"{compute_type}: {type(e).__name__}: {e}")
            _whisper = None
    raise RuntimeError("; ".join(errors))


def _transcribe_in_subprocess(path, timeout=900):
    """
    faster-whisper / ctranslate2 / PyAV can hard-crash the interpreter on Windows (no Python traceback,
    the whole run just exits). Run it in a child process so a crash only fails this tool call.
    """
    env = dict(os.environ)
    env.setdefault("KMP_DUPLICATE_LIB_OK", "TRUE")   # common OpenMP clash on Windows
    env.setdefault("OMP_NUM_THREADS", "4")
    env.setdefault("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")
    env["PYTHONIOENCODING"] = "utf-8"
    proc = subprocess.run(
        [sys.executable, "-X", "faulthandler", str(Path(__file__).resolve()), "--transcribe", str(path)],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout, env=env,
    )
    out = proc.stdout
    if proc.returncode == 0 and "<<<TRANSCRIPT>>>" in out:
        return out.split("<<<TRANSCRIPT>>>", 1)[1].strip()
    raise RuntimeError(f"whisper child process failed (exit code {proc.returncode}). "
                       f"stderr tail:\n{proc.stderr[-2500:]}")


@tool
def transcribe_audio(file_name: str) -> str:
    """
    Transcribe an audio file (mp3/wav/m4a/flac...) to text with timestamps, using local Whisper.

    Args:
        file_name: Attachment file name (or path).
    """
    try:
        path = resolve_file(file_name)
        cache = CACHE_DIR / f"{path.name}.{WHISPER_MODEL}.txt"
        if cache.exists():
            return cache.read_text(encoding="utf-8")
        text = _transcribe_in_subprocess(path)
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        cache.write_text(text, encoding="utf-8")
        return text
    except ImportError:
        return "faster-whisper not installed: pip install faster-whisper"
    except Exception as e:
        return f"transcribe_audio failed: {type(e).__name__}: {e}"


# =====================================================
# Vision (DeepSeek deepseek-flash accepts images)
# =====================================================

_vision_client = None


def _vision():
    global _vision_client
    if _vision_client is None:
        from openai import OpenAI

        _vision_client = OpenAI(api_key=os.getenv("DEEPSEEK_API_KEY"), base_url=API_BASE,
                                timeout=180, max_retries=1)
    return _vision_client


def _image_part(data_bytes, mime="image/png", detail="high"):
    b64 = base64.b64encode(data_bytes).decode()
    return {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}", "detail": detail}}


def vision_query(prompt, images, detail="high"):
    """images: list of (bytes, mime)."""
    content = [{"type": "text", "text": prompt}]
    content += [_image_part(b, m, detail) for b, m in images]
    resp = _vision().chat.completions.create(
        model=VISION_MODEL, messages=[{"role": "user", "content": content}]
    )
    return (resp.choices[0].message.content or "").strip()


@tool
def analyze_image(file_name: str, question: str) -> str:
    """
    Ask a vision model about an image file (screenshots, charts, tables, photos, chess boards...).
    Be specific: e.g. "List every piece and its square, files a-h from White's side".

    Args:
        file_name: Attachment file name (or path) of the image.
        question: What to look for / extract from the image.
    """
    try:
        path = resolve_file(file_name)
        mime = mimetypes.guess_type(str(path))[0] or "image/png"
        prompt = (
            "Look at the image carefully and answer precisely. Transcribe any text/numbers exactly. "
            "If something is uncertain, say so explicitly.\n\n" + question
        )
        return vision_query(prompt, [(path.read_bytes(), mime)])
    except Exception as e:
        return f"analyze_image failed: {type(e).__name__}: {e}"


# =====================================================
# YouTube
# =====================================================

def _youtube_id(url_or_id):
    m = re.search(r"(?:v=|youtu\.be/|shorts/|embed/)([A-Za-z0-9_-]{11})", url_or_id)
    if m:
        return m.group(1)
    if re.fullmatch(r"[A-Za-z0-9_-]{11}", url_or_id.strip()):
        return url_or_id.strip()
    raise ValueError(f"Cannot parse YouTube id from: {url_or_id}")


@tool
def youtube_transcript(url: str) -> str:
    """
    Get the transcript/captions (with timestamps) of a YouTube video. Good for "what does X say" questions.
    Captions do not show what is visible on screen; for visual questions use analyze_youtube_video.

    Args:
        url: YouTube URL or 11-character video id.
    """
    try:
        from youtube_transcript_api import YouTubeTranscriptApi

        vid = _youtube_id(url)
        api = YouTubeTranscriptApi()
        try:
            fetched = _run_with_timeout(lambda: api.fetch(vid, languages=["en", "en-US", "en-GB"]), 40)
        except Exception:
            tl = _run_with_timeout(lambda: api.list(vid), 40)
            fetched = next(iter(tl)).fetch()
        lines = [f"[{s.start:6.1f}s] {s.text}" for s in fetched]
        return _clip("\n".join(lines), 0, MAX_CHARS)
    except Exception as e:
        return (f"youtube_transcript failed: {type(e).__name__}: {str(e)[:300]}. "
                f"Try analyze_youtube_video, or web_search for the video title/quotes.")


def _download_youtube(vid):
    import yt_dlp

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    existing = sorted(CACHE_DIR.glob(f"yt_{vid}.*"))
    existing = [p for p in existing if p.suffix.lower() in (".mp4", ".webm", ".mkv") and p.stat().st_size > 0]
    if existing:
        return existing[0]
    base = {
        "outtmpl": str(CACHE_DIR / f"yt_{vid}.%(ext)s"),
        "quiet": True,
        "no_warnings": True,
        "noplaylist": True,
    }
    browser = os.getenv("YTDLP_COOKIES_FROM_BROWSER")  # e.g. "chrome" / "edge" if YouTube asks to sign in
    if browser:
        base["cookiesfrombrowser"] = (browser,)
    proxy = os.getenv("HTTPS_PROXY") or os.getenv("HTTP_PROXY")
    if proxy:
        base["proxy"] = proxy
    # We only need FRAMES -> a video-only stream is fine and needs no ffmpeg merge.
    # YouTube often no longer offers progressive (video+audio) mp4, which made "best[ext=mp4]" fail.
    formats = [
        "bv*[ext=mp4][vcodec^=avc1][height<=720]/bv*[ext=mp4][height<=720]",
        "bv*[height<=720]/bv*",
        "b[height<=720]/b",
        "18/worst",
    ]
    last = None
    for fmt in formats:
        try:
            with yt_dlp.YoutubeDL({**base, "format": fmt}) as ydl:
                info = ydl.extract_info(f"https://www.youtube.com/watch?v={vid}", download=True)
                path = Path(ydl.prepare_filename(info))
            if path.exists() and path.stat().st_size > 0:
                return path
        except Exception as e:
            last = e
    raise RuntimeError(f"yt-dlp could not download any format: {last}")


@tool
def analyze_youtube_video(url: str, question: str, max_frames: int = 48) -> str:
    """
    Download a YouTube video, sample frames evenly across it and ask a vision model the question.
    Use for visual questions (what is shown, how many objects/animals are on screen at once...).

    Args:
        url: YouTube URL or video id.
        question: The visual question to answer from the frames.
        max_frames: Number of frames to sample (default 48).
    """
    try:
        import cv2
    except ImportError:
        return "opencv not installed: pip install opencv-python-headless yt-dlp"
    try:
        vid = _youtube_id(url)
        path = _run_with_timeout(lambda: _download_youtube(vid), 240)
        cap = cv2.VideoCapture(str(path))
        n = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        fps = cap.get(cv2.CAP_PROP_FPS) or 25
        if n <= 0:
            return "Could not read video frames."
        idxs = sorted({int(i * (n - 1) / max(1, max_frames - 1)) for i in range(max_frames)})
        images, stamps = [], []
        for i in idxs:
            cap.set(cv2.CAP_PROP_POS_FRAMES, i)
            ok, frame = cap.read()
            if not ok:
                continue
            h, w = frame.shape[:2]
            scale = 768 / max(h, w)
            if scale < 1:
                frame = cv2.resize(frame, (int(w * scale), int(h * scale)))
            ok, buf = cv2.imencode(".jpg", frame, [cv2.IMWRITE_JPEG_QUALITY, 85])
            if ok:
                images.append((buf.tobytes(), "image/jpeg"))
                stamps.append(f"{i / fps:.1f}s")
        cap.release()
        prompt = (
            f"These are {len(images)} frames sampled in order from a video at times: {', '.join(stamps)}.\n"
            "Answer the question based on what is VISIBLE in the frames. "
            "Go frame by frame for counting questions and report the frame/time of your answer.\n\n"
            f"Question: {question}"
        )
        return vision_query(prompt, images, detail="high")
    except Exception as e:
        return (f"analyze_youtube_video failed: {type(e).__name__}: {str(e)[:300]}. "
                "(If YouTube is blocked in your network set HTTPS_PROXY.)")


# =====================================================
# Chess
# =====================================================

PIECE_CODES = set("KQRBNPkqrbnp.")


def _find_stockfish():
    p = os.getenv("STOCKFISH_PATH")
    if p and Path(p).exists():
        return p
    for name in ("stockfish", "stockfish.exe"):
        if shutil.which(name):
            return shutil.which(name)
    # any stockfish*.exe / stockfish* binary inside the project folder (e.g. unzipped release)
    for cand in sorted(BASE_DIR.glob("**/stockfish*")):
        if not cand.is_file() or ".venv" in cand.parts:
            continue
        if os.name == "nt":
            if cand.suffix.lower() == ".exe":
                return str(cand)
        elif cand.suffix == "" and os.access(cand, os.X_OK):
            return str(cand)
    return None


def _board_orientation(img_bytes):
    """Ask the vision model which side is at the bottom. Returns 'white' or 'black'."""
    ans = vision_query(
        "This is a chess diagram with coordinate labels (rank numbers 1-8 and file letters a-h) drawn "
        "on the edge squares. Look ONLY at the labels. What is the rank number written on the BOTTOM row, "
        "and what file letter is written on the LEFT-most column? Reply exactly in the form: "
        "bottom_rank=<number> left_file=<letter>",
        [(img_bytes, "image/png")], detail="high",
    )
    m = re.search(r"bottom_rank\s*=\s*(\d)", ans)
    if m:
        return "white" if m.group(1) == "1" else "black"
    m = re.search(r"left_file\s*=\s*([a-h])", ans)
    if m:
        return "white" if m.group(1) == "a" else "black"
    return None


def _classify_row(cells):
    """cells: list of 8 PNG bytes (one displayed row, left->right). Returns list of 8 codes."""
    prompt = (
        "You will see 8 small images, each is ONE square of a chess board, in order 1..8 (left to right). "
        "For each square tell which chess piece is on it.\n"
        "Codes: K king, Q queen, R rook, B bishop, N knight, P pawn. Use UPPERCASE for a WHITE/light-colored "
        "piece and lowercase for a BLACK/dark-colored piece. Use '.' for an empty square "
        "(a lone letter or digit label in the corner is NOT a piece).\n"
        "Hints: king has a cross on top; queen has a crown of several balls; bishop has a slit/mitre and one "
        "small ball; rook is a castle tower; knight is a horse head; pawn is a small round head.\n"
        'Reply with ONLY a JSON list of 8 strings, e.g. [".","p",".","K",".",".","b","."]'
    )
    ans = vision_query(prompt, [(c, "image/png") for c in cells], detail="high")
    m = re.search(r"\[.*?\]", ans, re.S)
    if not m:
        raise ValueError(f"unparseable: {ans[:200]}")
    codes = [str(x).strip() or "." for x in json.loads(m.group(0))]
    codes = [c if c in PIECE_CODES else "?" for c in codes]
    if len(codes) != 8:
        raise ValueError(f"expected 8 codes, got {codes}")
    return codes


@tool
def chess_image_to_fen(file_name: str, side_to_move: str) -> str:
    """
    Convert a chess board IMAGE into a FEN string reliably: the board is cut into 64 squares,
    every square is classified separately by the vision model (3 independent passes, majority vote),
    and board orientation (white or black at the bottom) is detected from the coordinate labels.
    Use this instead of describing the board with analyze_image. Then pass the FEN to chess_best_move.

    Args:
        file_name: Attachment file name (or path) of the board image.
        side_to_move: "w" or "b" (who is to move, from the question text).
    """
    try:
        from PIL import Image
    except ImportError:
        return "Pillow not installed: pip install pillow"
    try:
        path = resolve_file(file_name)
        img = Image.open(path).convert("RGB")
        W, H = img.size
        full_png = io.BytesIO()
        img.save(full_png, "PNG")
        full_bytes = full_png.getvalue()

        # 64 crops in DISPLAY order (row 0 = top)
        rows = []
        for r in range(8):
            row = []
            for c in range(8):
                box = (round(c * W / 8), round(r * H / 8), round((c + 1) * W / 8), round((r + 1) * H / 8))
                cell = img.crop(box).resize((160, 160))
                b = io.BytesIO()
                cell.save(b, "PNG")
                row.append(b.getvalue())
            rows.append(row)

        passes = 3
        jobs = [(p, r) for p in range(passes) for r in range(8)]
        results = {}
        with ThreadPoolExecutor(max_workers=12) as ex:
            futs = {ex.submit(_classify_row, rows[r]): (p, r) for p, r in jobs}
            orient_fut = ex.submit(_board_orientation, full_bytes)
            for f, key in futs.items():
                try:
                    results[key] = f.result(timeout=300)
                except Exception as e:
                    results[key] = ["?"] * 8
                    print(f"[chess] row {key} failed: {e}")
            orientation = orient_fut.result(timeout=300)

        from collections import Counter

        grid, uncertain = [], []
        for r in range(8):
            out_row = []
            for c in range(8):
                votes = Counter(results[(p, r)][c] for p in range(passes) if results[(p, r)][c] != "?")
                if not votes:
                    out_row.append(".")
                    uncertain.append((r, c, "no reading"))
                    continue
                code, n = votes.most_common(1)[0]
                if n < passes:
                    uncertain.append((r, c, dict(votes)))
                out_row.append(code)
            grid.append(out_row)

        # orientation fallback: where are the white pawns?
        if orientation is None:
            top = sum(row.count("P") for row in grid[:4])
            bottom = sum(row.count("P") for row in grid[4:])
            orientation = "black" if top > bottom else "white"

        def square_name(r, c):
            if orientation == "white":
                return "abcdefgh"[c] + str(8 - r)
            return "abcdefgh"[7 - c] + str(r + 1)

        # convert display grid -> standard grid (rank 8 first, file a first)
        std = grid if orientation == "white" else [row[::-1] for row in grid[::-1]]
        fen_rows = []
        for row in std:
            s, empty = "", 0
            for code in row:
                if code == ".":
                    empty += 1
                else:
                    if empty:
                        s += str(empty)
                        empty = 0
                    s += code
            if empty:
                s += str(empty)
            fen_rows.append(s)
        stm = "b" if side_to_move.strip().lower().startswith("b") else "w"
        fen = "/".join(fen_rows) + f" {stm} - - 0 1"

        report = [f"FEN: {fen}", f"orientation: {orientation} at the bottom of the image"]
        try:
            import chess

            board = chess.Board(fen)
            report.append(str(board))
            if not board.is_valid():
                report.append(f"WARNING: position not valid: {board.status()}")
        except Exception as e:
            report.append(f"(python-chess check failed: {e})")
        pieces = [f"{code}@{square_name(r, c)}" for r in range(8) for c in range(8) if (code := grid[r][c]) != "."]
        report.append("pieces: " + ", ".join(pieces))
        if uncertain:
            report.append("UNCERTAIN squares (votes differ): " + "; ".join(
                f"{square_name(r, c)} {v}" for r, c, v in uncertain))
            report.append("If needed, verify those squares with analyze_image before trusting the FEN.")
        return "\n".join(report)
    except Exception as e:
        return f"chess_image_to_fen failed: {type(e).__name__}: {e}"


_PIECE_VAL = {1: 100, 2: 320, 3: 330, 4: 500, 5: 900, 6: 0}


def _eval(board):
    import chess

    s = 0
    for pt, v in _PIECE_VAL.items():
        s += v * (len(board.pieces(pt, chess.WHITE)) - len(board.pieces(pt, chess.BLACK)))
    return s if board.turn == chess.WHITE else -s


class _SearchTimeout(Exception):
    pass


def _search(board, depth, alpha, beta, deadline, ply=0):
    if time.time() > deadline:
        raise _SearchTimeout()
    if board.is_checkmate():
        return -100000 + ply, None
    if board.is_stalemate() or board.is_insufficient_material() or board.can_claim_threefold_repetition():
        return 0, None
    if depth <= 0:
        return _quiesce(board, alpha, beta, deadline, 0), None
    best_move = None
    moves = sorted(board.legal_moves,
                   key=lambda m: (not board.gives_check(m), not board.is_capture(m)))
    for m in moves:
        board.push(m)
        score, _ = _search(board, depth - 1, -beta, -alpha, deadline, ply + 1)
        score = -score
        board.pop()
        if best_move is None or score > alpha:
            alpha, best_move = max(alpha, score), m
        if alpha >= beta:
            break
    return alpha, best_move


def _quiesce(board, alpha, beta, deadline, qd):
    stand = _eval(board)
    if stand >= beta or qd > 4:
        return stand
    alpha = max(alpha, stand)
    for m in board.legal_moves:
        if not board.is_capture(m):
            continue
        board.push(m)
        score = -_quiesce(board, -beta, -alpha, deadline, qd + 1)
        board.pop()
        if score >= beta:
            return score
        alpha = max(alpha, score)
    return alpha


@tool
def chess_best_move(fen: str) -> str:
    """
    Compute the best move for the side to move in a chess position given as FEN.
    Uses Stockfish if available (STOCKFISH_PATH, PATH, or a stockfish folder inside the project),
    otherwise a small built-in search (much weaker). Returns the move in SAN (e.g. "Nf3"), UCI and the line.

    Args:
        fen: Full FEN string, e.g. "r1b1k2r/pp3ppp/... b - - 0 1". Make sure the side to move is correct.
    """
    try:
        import chess
        import chess.engine
    except ImportError:
        return "python-chess not installed: pip install chess"
    try:
        board = chess.Board(fen)
    except Exception as e:
        return f"Invalid FEN: {e}"
    if not board.is_valid():
        return f"FEN is not a legal position ({board.status()}). Re-check the piece placement."

    engine_path = _find_stockfish()
    if engine_path:
        try:
            with chess.engine.SimpleEngine.popen_uci(engine_path) as eng:
                info = eng.analyse(board, chess.engine.Limit(time=5.0), multipv=3)
            lines = []
            for i, inf in enumerate(info, 1):
                pv = inf.get("pv", [])
                if not pv:
                    continue
                lines.append(f"{i}. {board.san(pv[0])} ({pv[0].uci()}) score={inf['score'].pov(board.turn)}"
                             f"  line: {board.variation_san(pv[:8])}")
            top = info[0]["pv"][0] if info and info[0].get("pv") else None
            head = (f"BEST MOVE (Stockfish, reliable): {board.san(top)}\n"
                    "Use this as the answer. No further verification is needed.\n") if top else ""
            return head + "Top lines:\n" + "\n".join(lines)
        except Exception as e:
            print(f"[tools] stockfish failed, fallback search: {e}")

    deadline = time.time() + 60
    best = None
    for depth in range(1, 8):
        try:
            score, move = _search(board.copy(), depth, -10**9, 10**9, deadline)
            if move is not None:
                best = (depth, score, move)
            if abs(score) > 90000:  # forced mate found
                break
        except _SearchTimeout:
            break
    if not best:
        return "No legal move found."
    d, score, move = best
    note = "forced mate" if score > 90000 else f"material score {score}"
    return (f"Built-in search (depth {d}, {note}): best move {board.san(move)} ({move.uci()}). "
            "This fallback is weak; install Stockfish (see README) for a reliable answer.")


# =====================================================
# Export
# =====================================================

ALL_TOOLS = [
    web_search,
    visit_webpage,
    wikipedia_search,
    wikipedia_page,
    read_file,
    run_python_file,
    transcribe_audio,
    analyze_image,
    youtube_transcript,
    analyze_youtube_video,
    chess_image_to_fen,
    chess_best_move,
]


# =====================================================
# Child-process entry point (used by transcribe_audio)
#   python tools.py --transcribe path/to/file.mp3
# =====================================================

if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--transcribe":
        result = transcribe_audio_path(sys.argv[2])
        sys.stdout.write("<<<TRANSCRIPT>>>\n" + result + "\n")
        sys.stdout.flush()
